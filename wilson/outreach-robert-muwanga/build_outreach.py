# -*- coding: utf-8 -*-
"""
Build the outreach package for the introduction to Mr. Robert Muwanga.

    python build_outreach.py

Produces
--------
  email_to_robert_muwanga.txt        the email, ready to paste into a mail client
  Wasswa_Wilson_Cover_Letter.pdf     one page, matching the CV house style
  Wasswa_Wilson_Cover_Letter.docx    editable version of the same
  Wasswa_Wilson_Dossier_No_ID.pdf    contents + letter + CV + credentials
  Wasswa_Wilson_Full_Dossier.pdf     the same, plus identification documents

The CV itself is not rebuilt here. It is taken from ../cv-2026/, so the CV
stays the single source of truth for its own content.
"""

import io
import os
import sys
import zipfile

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import fitz
from PIL import Image

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.platypus import (BaseDocTemplate, Frame, Image as RLImage,
                                KeepTogether, PageTemplate, Paragraph, Spacer,
                                Table, TableStyle)

import letter_content as L

HERE = os.path.dirname(os.path.abspath(__file__))
WILSON = os.path.dirname(HERE)
CV_PDF = os.path.join(WILSON, "cv-2026", "Wasswa_Wilson_CV_2026.pdf")
SUPPORT_DOCX = os.path.join(WILSON, "WASSWA WILSON - SUPPORTING DOCUMENTS.docx")
ID_PDF = os.path.join(WILSON, "Identification documents.pdf")
SIGNATURE = os.path.join(WILSON, "medequip-rwanda", "signature_clean.png")

EMAIL_OUT = os.path.join(HERE, "email_to_robert_muwanga.txt")
LETTER_PDF = os.path.join(HERE, "Wasswa_Wilson_Cover_Letter.pdf")
LETTER_DOCX = os.path.join(HERE, "Wasswa_Wilson_Cover_Letter.docx")
DOSSIER_FULL = os.path.join(HERE, "Wasswa_Wilson_Full_Dossier.pdf")
DOSSIER_NOID = os.path.join(HERE, "Wasswa_Wilson_Dossier_No_ID.pdf")

# ---------------------------------------------------------------- palette --
NAVY_HEX, TEAL_HEX = "0F2C4C", "1A6B5C"
PALE_HEX, INK_HEX = "9FD6C8", "23272B"
MUTED_HEX, HAIR_HEX, BAND_HEX = "5A6472", "CFD8E3", "F3F6F9"

NAVY = colors.HexColor("#" + NAVY_HEX)
TEAL = colors.HexColor("#" + TEAL_HEX)
PALE = colors.HexColor("#" + PALE_HEX)
INK = colors.HexColor("#" + INK_HEX)
MUTED = colors.HexColor("#" + MUTED_HEX)
HAIR = colors.HexColor("#" + HAIR_HEX)
BAND = colors.HexColor("#" + BAND_HEX)

NAVY_RGB = RGBColor(0x0F, 0x2C, 0x4C)
TEAL_RGB = RGBColor(0x1A, 0x6B, 0x5C)
INK_RGB = RGBColor(0x23, 0x27, 0x2B)
MUTED_RGB = RGBColor(0x5A, 0x64, 0x72)
WHITE_RGB = RGBColor(0xFF, 0xFF, 0xFF)

PW, PH = A4
LM = RM = 0.72 * inch
TM, BM = 0.5 * inch, 0.6 * inch
CW = PW - LM - RM


# ==========================================================================
#  1. Email
# ==========================================================================

def build_email():
    text = ("Subject: %s\n\n%s" % (L.EMAIL_SUBJECT, L.EMAIL_BODY))
    with io.open(EMAIL_OUT, "w", encoding="utf-8") as fh:
        fh.write(text)
    return EMAIL_OUT


# ==========================================================================
#  2. Cover letter, PDF
# ==========================================================================

def letter_styles():
    s = {}
    s["name"] = ParagraphStyle("n", fontName="Helvetica-Bold", fontSize=20, leading=23,
                               textColor=colors.white, alignment=1, spaceAfter=3)
    s["role"] = ParagraphStyle("r", fontName="Helvetica-Bold", fontSize=9, leading=11.5,
                               textColor=PALE, alignment=1)
    s["contact"] = ParagraphStyle("c", fontName="Helvetica", fontSize=8.8, leading=11.5,
                                  textColor=MUTED, alignment=1)
    s["meta"] = ParagraphStyle("m", fontName="Helvetica", fontSize=9.2, leading=12.6,
                               textColor=INK)
    s["date"] = ParagraphStyle("d", fontName="Helvetica", fontSize=9.2, leading=12.6,
                               textColor=MUTED)
    s["subject"] = ParagraphStyle("s", fontName="Helvetica-Bold", fontSize=9.8, leading=13,
                                  textColor=NAVY)
    s["h"] = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=9.5, leading=12.6,
                            textColor=TEAL, spaceBefore=4, spaceAfter=2)
    s["p"] = ParagraphStyle("p", fontName="Helvetica", fontSize=9.5, leading=12.6,
                            textColor=INK, alignment=TA_JUSTIFY, spaceAfter=6)
    s["sig"] = ParagraphStyle("sg", fontName="Helvetica-Bold", fontSize=9.8, leading=12.5,
                              textColor=NAVY)
    s["sigsub"] = ParagraphStyle("ss", fontName="Helvetica", fontSize=8.8, leading=11.5,
                                 textColor=MUTED)
    return s


def build_letter_pdf():
    s = letter_styles()
    doc = BaseDocTemplate(LETTER_PDF, pagesize=A4, leftMargin=LM, rightMargin=RM,
                          topMargin=TM, bottomMargin=BM,
                          title="Wasswa Wilson, letter of introduction",
                          author="Wasswa Wilson")
    doc.addPageTemplates([PageTemplate(id="all", frames=[
        Frame(LM, BM, CW, PH - TM - BM, id="b",
              leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)])])

    st = []
    mast = Table([[[Paragraph("&nbsp;".join(L.SENDER["name"].upper()), s["name"]),
                    Paragraph(L.SENDER["title"], s["role"])]]], colWidths=[CW])
    mast.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), NAVY),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 11),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 11),
        ("LINEBELOW", (0, 0), (-1, -1), 3.0, TEAL)]))
    st.append(mast)
    st.append(Spacer(1, 6))
    st.append(Paragraph("%s  ·  %s  ·  %s" % (L.SENDER["city"], L.SENDER["email"],
                                              L.SENDER["phone"]), s["contact"]))
    st.append(Spacer(1, 16))

    st.append(Paragraph(L.DATE, s["date"]))
    st.append(Spacer(1, 9))
    for line in L.RECIPIENT:
        st.append(Paragraph(line, s["meta"]))
    st.append(Spacer(1, 13))

    subj = Table([[Paragraph(L.SUBJECT_LINE, s["subject"])]], colWidths=[CW])
    subj.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), BAND),
        ("LINEBEFORE", (0, 0), (0, 0), 2.6, TEAL),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    st.append(subj)
    st.append(Spacer(1, 13))

    st.append(Paragraph(L.SALUTATION, s["p"]))
    for heading, text in L.LETTER_BODY:
        if heading:
            st.append(Paragraph(heading, s["h"]))
        st.append(Paragraph(text, s["p"]))

    st.append(Spacer(1, 4))
    sig_block = [Paragraph(L.CLOSING, s["p"])]
    if os.path.exists(SIGNATURE):
        img = Image.open(SIGNATURE)
        h = 0.42 * inch
        w = h * img.size[0] / float(img.size[1])
        sig_block.append(RLImage(SIGNATURE, width=w, height=h, hAlign="LEFT"))
        sig_block.append(Spacer(1, 3))
    sig_block += [Paragraph(L.SENDER["name"], s["sig"]),
                  Paragraph(L.SENDER["title"], s["sigsub"]),
                  Paragraph("%s  ·  %s" % (L.SENDER["email"], L.SENDER["phone"]), s["sigsub"])]
    st.append(KeepTogether(sig_block))

    doc.build(st)
    return LETTER_PDF


# ==========================================================================
#  3. Cover letter, DOCX
# ==========================================================================

def _cell_shade(cell, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _cell_border(cell, edge, hex_color, size):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    e = OxmlElement("w:" + edge)
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(size))
    e.set(qn("w:space"), "0")
    e.set(qn("w:color"), hex_color)
    b.append(e)


def _no_borders(table):
    bs = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        bs.append(e)
    table._tbl.tblPr.append(bs)


def _margins(table, top=0, left=0, bottom=0, right=0):
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    table._tbl.tblPr.append(mar)


def _run(par, text, size=10, bold=False, italic=False, color=INK_RGB, spacing=None):
    r = par.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    rPr = r._element.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), "Calibri")
    rPr.append(rf)
    if spacing:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        rPr.append(sp)
    return r


def _p(doc, before=0, after=6, align=None, line=1.06):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def build_letter_docx():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    sec.left_margin = sec.right_margin = Inches(0.72)
    content_w = 8.27 - 1.44

    mast = doc.add_table(rows=1, cols=1)
    mast.autofit = False
    _no_borders(mast)
    _margins(mast, top=160, left=120, bottom=160, right=120)
    cell = mast.rows[0].cells[0]
    cell.width = Inches(content_w)
    _cell_shade(cell, NAVY_HEX)
    _cell_border(cell, "bottom", TEAL_HEX, 24)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    _run(p, L.SENDER["name"].upper(), size=20, bold=True, color=WHITE_RGB, spacing=44)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    _run(p, L.SENDER["title"], size=9, bold=True, color=RGBColor(0x9F, 0xD6, 0xC8))

    p = _p(doc, before=6, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, "%s  ·  %s  ·  %s" % (L.SENDER["city"], L.SENDER["email"], L.SENDER["phone"]),
         size=8.8, color=MUTED_RGB)

    p = _p(doc, after=8)
    _run(p, L.DATE, size=9.2, color=MUTED_RGB)
    for line in L.RECIPIENT:
        p = _p(doc, after=1)
        _run(p, line, size=9.2)

    subj = doc.add_table(rows=1, cols=1)
    subj.autofit = False
    _no_borders(subj)
    _margins(subj, top=70, left=110, bottom=70, right=110)
    c = subj.rows[0].cells[0]
    c.width = Inches(content_w)
    _cell_shade(c, BAND_HEX)
    _cell_border(c, "left", TEAL_HEX, 24)
    sp = c.paragraphs[0]
    sp.paragraph_format.space_after = Pt(0)
    _run(sp, L.SUBJECT_LINE, size=9.8, bold=True, color=NAVY_RGB)

    p = _p(doc, before=12, after=7)
    _run(p, L.SALUTATION, size=9.5)

    for heading, text in L.LETTER_BODY:
        if heading:
            hp = _p(doc, before=4, after=2)
            hp.paragraph_format.keep_with_next = True
            _run(hp, heading, size=9.5, bold=True, color=TEAL_RGB)
        tp = _p(doc, after=7, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        _run(tp, text, size=9.5)

    p = _p(doc, before=4, after=4)
    _run(p, L.CLOSING, size=9.5)
    if os.path.exists(SIGNATURE):
        sp = _p(doc, after=2)
        sp.add_run().add_picture(SIGNATURE, height=Inches(0.42))
    p = _p(doc, after=1)
    _run(p, L.SENDER["name"], size=9.8, bold=True, color=NAVY_RGB)
    p = _p(doc, after=1)
    _run(p, L.SENDER["title"], size=8.8, color=MUTED_RGB)
    p = _p(doc, after=0)
    _run(p, "%s  ·  %s" % (L.SENDER["email"], L.SENDER["phone"]), size=8.8, color=MUTED_RGB)

    doc.save(LETTER_DOCX)
    return LETTER_DOCX


# ==========================================================================
#  4. Contents page and captioned document pages
# ==========================================================================

def draw_footer(c, text):
    c.setFont("Helvetica", 7.8)
    c.setFillColor(MUTED)
    c.drawCentredString(PW / 2.0, 0.32 * inch, text)


def draw_caption(c, section_label, caption, index, total):
    """Teal bar + navy section label + grey caption at the top of a document page."""
    top = PH - 0.5 * inch
    c.setFillColor(TEAL)
    c.rect(LM, top - 13, 2.6, 13, stroke=0, fill=1)
    c.setFont("Helvetica-Bold", 9.4)
    c.setFillColor(NAVY)
    c.drawString(LM + 8, top - 10, section_label.upper())
    c.setFont("Helvetica", 8.2)
    c.setFillColor(MUTED)
    c.drawRightString(PW - RM, top - 10, "%d of %d" % (index, total))
    c.setStrokeColor(HAIR)
    c.setLineWidth(0.5)
    c.line(LM, top - 17, PW - RM, top - 17)
    # caption, wrapped by hand at the available width
    c.setFont("Helvetica", 8.6)
    c.setFillColor(INK)
    words, line, y = caption.split(), "", top - 29
    avail = CW
    for w in words:
        trial = (line + " " + w).strip()
        if c.stringWidth(trial, "Helvetica", 8.6) > avail:
            c.drawString(LM, y, line)
            y -= 10.5
            line = w
        else:
            line = trial
    if line:
        c.drawString(LM, y, line)
    return y - 10


def place_image(c, pil_img, top_y):
    """Scale a PIL image to fit the space between top_y and the footer."""
    avail_h = top_y - 0.55 * inch
    avail_w = CW
    iw, ih = pil_img.size
    scale = min(avail_w / float(iw), avail_h / float(ih))
    w, h = iw * scale, ih * scale
    x = LM + (avail_w - w) / 2.0
    y = 0.55 * inch + (avail_h - h) / 2.0
    buf = io.BytesIO()
    pil_img.save(buf, format="JPEG", quality=82, optimize=True)
    buf.seek(0)
    c.drawImage(ImageReader(buf), x, y, width=w, height=h)
    # light frame so scans sit on the page deliberately
    c.setStrokeColor(HAIR)
    c.setLineWidth(0.5)
    c.rect(x, y, w, h, stroke=1, fill=0)


def load_credentials():
    z = zipfile.ZipFile(SUPPORT_DOCX)
    out = []
    for fname, caption in L.CREDENTIALS:
        im = Image.open(io.BytesIO(z.read("word/media/" + fname))).convert("RGB")
        if fname in L.ROTATE_CCW:
            im = im.transpose(Image.ROTATE_90)
        im.thumbnail((1700, 1700), Image.LANCZOS)
        out.append((im, caption))
    return out


def load_identification():
    d = fitz.open(ID_PDF)
    out = []
    for idx, caption in L.IDENTIFICATION:
        pix = d[idx].get_pixmap(dpi=200)
        im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        im.thumbnail((1700, 1700), Image.LANCZOS)
        out.append((im, caption))
    d.close()
    return out


def build_front_matter(path, contents, n_pages_note):
    """Contents page, as its own single-page PDF."""
    c = pdfcanvas.Canvas(path, pagesize=A4)
    c.setFillColor(NAVY)
    c.rect(0, PH - 2.35 * inch, PW, 2.35 * inch, stroke=0, fill=1)
    c.setFillColor(TEAL)
    c.rect(0, PH - 2.35 * inch - 4, PW, 4, stroke=0, fill=1)

    c.setFont("Helvetica-Bold", 26)
    c.setFillColor(colors.white)
    name = " ".join(L.SENDER["name"].upper())
    c.drawCentredString(PW / 2.0, PH - 1.28 * inch, name)
    c.setFont("Helvetica-Bold", 10.5)
    c.setFillColor(PALE)
    c.drawCentredString(PW / 2.0, PH - 1.58 * inch, L.SENDER["title"])
    c.setFont("Helvetica", 9.4)
    c.drawCentredString(PW / 2.0, PH - 1.86 * inch,
                        "%s  ·  %s  ·  %s" % (L.SENDER["city"], L.SENDER["email"],
                                              L.SENDER["phone"]))

    y = PH - 3.1 * inch
    c.setFont("Helvetica-Bold", 13)
    c.setFillColor(NAVY)
    c.drawString(LM, y, L.DOSSIER_TITLE.upper())
    c.setStrokeColor(TEAL)
    c.setLineWidth(2)
    c.line(LM, y - 8, LM + 1.5 * inch, y - 8)

    y -= 34
    c.setFont("Helvetica", 9.6)
    c.setFillColor(MUTED)
    c.drawString(LM, y, "Prepared for Mr. Robert Muwanga  ·  %s" % L.DATE)

    y -= 30
    for num, title, note in contents:
        c.setFillColor(BAND)
        c.rect(LM, y - 17, CW, 30, stroke=0, fill=1)
        c.setFillColor(TEAL)
        c.rect(LM, y - 17, 2.6, 30, stroke=0, fill=1)
        c.setFont("Helvetica-Bold", 11)
        c.setFillColor(TEAL)
        c.drawString(LM + 12, y + 2, num)
        c.setFont("Helvetica-Bold", 10.4)
        c.setFillColor(NAVY)
        c.drawString(LM + 30, y + 2, title)
        c.setFont("Helvetica", 8.8)
        c.setFillColor(MUTED)
        c.drawString(LM + 30, y - 10, note)
        y -= 40

    y -= 6
    c.setFont("Helvetica-Oblique", 8.6)
    c.setFillColor(MUTED)
    c.drawString(LM, y, n_pages_note)

    draw_footer(c, "Wasswa Wilson  ·  Professional Dossier")
    c.showPage()
    c.save()
    return path


def build_docs_pdf(path, creds, ids):
    c = pdfcanvas.Canvas(path, pagesize=A4)
    total = len(creds)
    for i, (im, caption) in enumerate(creds, 1):
        y = draw_caption(c, "Academic and professional credentials", caption, i, total)
        place_image(c, im, y)
        draw_footer(c, "Wasswa Wilson  ·  Supporting Documents")
        c.showPage()
    if ids:
        total = len(ids)
        for i, (im, caption) in enumerate(ids, 1):
            y = draw_caption(c, "Identification documents", caption, i, total)
            place_image(c, im, y)
            draw_footer(c, "Wasswa Wilson  ·  Supporting Documents")
            c.showPage()
    c.save()
    return path


# ==========================================================================
#  5. Assembly
# ==========================================================================

def assemble(out_path, contents, creds, ids, note):
    tmp_front = out_path + ".front.tmp.pdf"
    tmp_docs = out_path + ".docs.tmp.pdf"
    build_front_matter(tmp_front, contents, note)
    build_docs_pdf(tmp_docs, creds, ids)

    merged = fitz.open()
    for part in (tmp_front, LETTER_PDF, CV_PDF, tmp_docs):
        with fitz.open(part) as src:
            merged.insert_pdf(src)
    merged.set_metadata({
        "title": "Wasswa Wilson, Professional Dossier",
        "author": "Wasswa Wilson",
        "subject": "Curriculum vitae and supporting documents",
    })
    merged.save(out_path, garbage=4, deflate=True)
    n = merged.page_count
    merged.close()
    for t in (tmp_front, tmp_docs):
        os.remove(t)
    return out_path, n


if __name__ == "__main__":
    print("email   ->", build_email())
    print("letter  ->", build_letter_pdf())
    print("letter  ->", build_letter_docx())

    creds = load_credentials()
    ids = load_identification()

    p, n = assemble(DOSSIER_NOID, L.CONTENTS_NO_ID, creds, [],
                    "This file contains the cover letter, the curriculum vitae and the "
                    "academic and professional credentials.")
    print("dossier ->", p, "(%d pages, %.1f MB)" % (n, os.path.getsize(p) / 1e6))

    p, n = assemble(DOSSIER_FULL, L.CONTENTS, creds, ids,
                    "This file contains the cover letter, the curriculum vitae, the academic "
                    "and professional credentials and the identification documents.")
    print("dossier ->", p, "(%d pages, %.1f MB)" % (n, os.path.getsize(p) / 1e6))
